"""Command-line interface for reasoning kits."""

import argparse
import asyncio
import re
import sys
from pathlib import Path

from dotenv import load_dotenv

from .loader import list_reasoning_kits, load_reasoning_kit


def resolve_kit_path(kit: str, base_path: str) -> Path:
    """Resolve a kit name to its directory path."""
    return Path(base_path) / kit


def main() -> None:
    """Main entry point for the CLI."""
    load_dotenv()

    parser = argparse.ArgumentParser(
        description="Reasoning Kits — a lightweight standard for self-improving LLM workflows.",
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    subparsers = parser.add_subparsers(dest="command", help="Available commands")

    list_parser = subparsers.add_parser("list", help="List available reasoning kits")
    list_parser.add_argument("--path", type=str, default="kits", help="Base path for local reasoning kits")

    info_parser = subparsers.add_parser("info", help="Show info about a reasoning kit")
    info_parser.add_argument("kit", type=str)
    info_parser.add_argument("--base-path", type=str, default="kits")

    validate_parser = subparsers.add_parser("validate", help="Validate a kit structure")
    validate_parser.add_argument("kit", type=str)
    validate_parser.add_argument("--base-path", type=str, default="kits")

    run_parser = subparsers.add_parser("run", help="Run a reasoning kit")
    run_parser.add_argument("kit", type=str)
    run_parser.add_argument("--base-path", type=str, default="kits")
    run_parser.add_argument("--evaluate", action="store_true", help="Enable step-by-step evaluation")
    run_parser.add_argument("--mode", choices=["transparent", "anonymous"], default=None)
    run_parser.add_argument("--model", type=str, default=None)
    run_parser.add_argument("--auto-evaluate", action="store_true", help="Score steps via judge LLM")
    run_parser.add_argument("--judge-model", type=str, default=None)
    run_parser.add_argument(
        "--evaluator-aggregation",
        choices=["average", "min", "max", "first"],
        default=None,
    )
    run_parser.add_argument("--dynamic-resource", action="append", default=[])
    run_parser.add_argument("--dynamic-resource-file", action="append", default=[])
    run_parser.add_argument("--stdin", action="append", default=[])
    run_parser.add_argument("--verbose", "-v", action="store_true")
    run_parser.add_argument("--no-mcp", action="store_true", help="Skip MCP initialization")
    run_parser.add_argument("--output-dir", type=str, default="outputs")
    run_parser.add_argument("--no-save", action="store_true")

    batch_parser = subparsers.add_parser("batch", help="Run a kit over multiple input files")
    batch_parser.add_argument("kit", type=str)
    batch_parser.add_argument("--input-dir", type=str, required=True)
    batch_parser.add_argument("--resource", type=str, required=True)
    batch_parser.add_argument("--dynamic", action="append", default=[])
    batch_parser.add_argument("--model", type=str, default=None)
    batch_parser.add_argument("--base-path", type=str, default="kits")
    batch_parser.add_argument("--output-dir", type=str, default="outputs")
    batch_parser.add_argument("--no-save", action="store_true")
    batch_parser.add_argument("--verbose", "-v", action="store_true")

    args = parser.parse_args()

    if args.command == "list":
        _cmd_list(args)
    elif args.command == "info":
        _cmd_info(args)
    elif args.command == "validate":
        _cmd_validate(args)
    elif args.command == "run":
        _cmd_run(args)
    elif args.command == "batch":
        _cmd_batch(args)
    else:
        parser.print_help()


def _cmd_list(args: argparse.Namespace) -> None:
    kits = list_reasoning_kits(args.path)
    if kits:
        print("Available reasoning kits:")
        for kit in kits:
            print(f"  - {kit}")
    else:
        print(f"No reasoning kits found in {args.path}")


def _cmd_info(args: argparse.Namespace) -> None:
    try:
        kit_path = resolve_kit_path(args.kit, args.base_path)
        kit = load_reasoning_kit(kit_path)
    except FileNotFoundError as e:
        print(f"Error: {e}")
        sys.exit(1)

    print(f"\nReasoning Kit: {kit.name}")
    print(f"Path: {kit.path}")
    print(f"\nResources ({len(kit.resources)}):")
    for key, resource in kit.resources.items():
        print(f"  {key}. {resource.resource_id} ({resource.file})")
    print(f"\nWorkflow Steps ({len(kit.workflow)}):")
    for key, step in kit.workflow.items():
        print(f"  {key}. {step.output_id} ({step.file})")
    if kit.tools:
        print(f"\nTools ({len(kit.tools)}):")
        for key, tool in kit.tools.items():
            display = f" ({tool.display_name})" if tool.display_name else ""
            print(f"  {key}. {tool.tool_name}{display}")
    if kit.evaluators:
        print(f"\nEvaluators ({len(kit.evaluators)} step(s)):")
        for step_key in sorted(kit.evaluators.keys(), key=int):
            pairs = kit.evaluators[step_key]
            labels = ", ".join(label for label, _ in pairs)
            dim_str = f"({len(pairs)} dimension{'s' if len(pairs) != 1 else ''})"
            print(f"  Step {step_key}: {labels}  {dim_str}")


def _cmd_validate(args: argparse.Namespace) -> None:
    from .tools import get_tool

    errors: list[str] = []
    warnings: list[str] = []

    try:
        kit_path = resolve_kit_path(args.kit, args.base_path)
        kit = load_reasoning_kit(kit_path)
    except FileNotFoundError as e:
        print(f"Error: {e}")
        sys.exit(1)

    step_nums = sorted(int(k) for k in kit.workflow.keys())
    expected_steps = list(range(1, len(step_nums) + 1))
    if step_nums != expected_steps:
        errors.append(
            f"Instructions are not numbered sequentially. Found: {step_nums}, expected: {expected_steps}"
        )

    resource_ids = {r.resource_id for r in kit.resources.values()}
    for step_num, step in kit.workflow.items():
        placeholders = re.findall(r"\{(resource_\d+)\}", step.prompt)
        for ph in placeholders:
            if ph not in resource_ids:
                errors.append(f"Step {step_num} references missing resource {{{ph}}}")

    tool_ids = {t.tool_id for t in kit.tools.values()}
    for step_num, step in kit.workflow.items():
        placeholders = re.findall(r"\{(tool_\d+)\}", step.prompt)
        for ph in placeholders:
            if ph not in tool_ids:
                errors.append(f"Step {step_num} references missing tool {{{ph}}}")
            else:
                tool_num = ph.replace("tool_", "")
                tool = kit.tools[tool_num]
                if tool.tool_name and get_tool(tool.tool_name) is None:
                    warnings.append(
                        f"Step {step_num} references tool '{tool.tool_name}' which is not in the global registry"
                    )

    for step_num, step in kit.workflow.items():
        placeholders = re.findall(r"\{(workflow_(\d+))\}", step.prompt)
        for full_ph, ref_num_str in placeholders:
            if int(ref_num_str) >= int(step_num):
                errors.append(f"Step {step_num} references {{{full_ph}}} which is not an earlier step")

    print(f"\nValidation Results for '{kit.name}':")
    print("=" * 50)
    if errors:
        print(f"\nErrors ({len(errors)}):")
        for err in errors:
            print(f"  ❌ {err}")
    if warnings:
        print(f"\nWarnings ({len(warnings)}):")
        for warn in warnings:
            print(f"  ⚠️  {warn}")
    if not errors and not warnings:
        print("\n✅ Kit is valid!")
    print()
    if errors:
        sys.exit(1)


def _cmd_run(args: argparse.Namespace) -> None:
    from .mcp_client import close_mcp_servers, init_mcp_servers
    from .runner import run_kit_async

    try:
        kit_path = resolve_kit_path(args.kit, args.base_path)
        local_mcp = kit_path / "mcp_servers.json"
        kit_config = str(local_mcp) if local_mcp.exists() else None

        kit = load_reasoning_kit(kit_path)

        dynamic_inputs: dict[str, str] = {}
        for stdin_spec in args.stdin:
            dynamic_inputs[stdin_spec] = sys.stdin.read()

        for file_spec in args.dynamic_resource_file:
            if "=" not in file_spec:
                print(f"Error: Invalid --dynamic-resource-file format: {file_spec}")
                sys.exit(1)
            res_id, file_path = file_spec.split("=", 1)
            path = Path(file_path)
            if not path.exists():
                print(f"Error: File not found for dynamic resource: {path}")
                sys.exit(1)
            dynamic_inputs[res_id] = path.read_text()

        for inline_spec in args.dynamic_resource:
            if "=" not in inline_spec:
                print(f"Error: Invalid --dynamic-resource format: {inline_spec}")
                sys.exit(1)
            res_id, text = inline_spec.split("=", 1)
            dynamic_inputs[res_id] = text

        dynamic_resources = [r for r in kit.resources.values() if r.is_dynamic]
        for resource in dynamic_resources:
            if resource.resource_id in dynamic_inputs:
                resource.content = dynamic_inputs[resource.resource_id]
            else:
                while True:
                    file_path = input(
                        f"  Enter file path for {resource.resource_id} ({resource.file}): "
                    ).strip()
                    if not file_path:
                        print("    File path cannot be empty. Please try again.")
                        continue
                    path = Path(file_path)
                    if not path.exists():
                        print(f"    File not found: {path}. Please try again.")
                        continue
                    try:
                        resource.content = path.read_text()
                        print(f"    Loaded {len(resource.content)} characters from {path.name}")
                        break
                    except Exception as e:
                        print(f"    Error reading file: {e}. Please try again.")

        collected_prompts: dict[str, str] = {}

        cfg = kit.config
        effective_mode = args.mode or cfg.mode or "transparent"
        effective_aggregation = args.evaluator_aggregation or cfg.evaluator_aggregation or "average"
        effective_judge_model = args.judge_model or cfg.judge_model
        effective_model = args.model or cfg.model

        async def _run_with_mcp() -> dict[str, str]:
            if not args.no_mcp:
                await init_mcp_servers(
                    config_path="mcp_servers.json",
                    kit_config_path=kit_config,
                )
            try:
                evaluate = args.evaluate or args.auto_evaluate
                kwargs: dict = dict(
                    evaluate=evaluate,
                    evaluation_mode=effective_mode,
                    verbose=args.verbose,
                    collected_prompts=collected_prompts,
                    auto_evaluate=args.auto_evaluate,
                    judge_model=effective_judge_model,
                    evaluator_aggregation=effective_aggregation,
                )
                if effective_model:
                    kwargs["model"] = effective_model
                return await run_kit_async(kit, **kwargs)
            finally:
                if not args.no_mcp:
                    await close_mcp_servers()

        outputs = asyncio.run(_run_with_mcp())

        if not args.no_save:
            from datetime import datetime

            timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
            run_path = Path(args.output_dir) / kit.name / timestamp
            run_path.mkdir(parents=True, exist_ok=True)
            _write_run_files(kit, outputs, collected_prompts, run_path)
            print(f"\nOutputs saved to: {run_path}")
    except FileNotFoundError as e:
        print(f"Error: {e}")
        sys.exit(1)
    except Exception as e:
        print(f"Error running reasoning kit: {e}")
        sys.exit(1)


def _write_run_files(kit, outputs: dict[str, str], prompts: dict[str, str], run_dir: Path) -> None:
    """Write output and prompt files for a completed kit run into run_dir."""
    import json
    from datetime import datetime

    step_meta = []
    for step_key in sorted(kit.workflow.keys(), key=int):
        step = kit.workflow[step_key]
        step_num = int(step_key)
        base = f"step_{step_num}_{step.output_id}"

        output_file = f"{base}.md"
        (run_dir / output_file).write_text(outputs.get(step.output_id, ""), encoding="utf-8")

        prompt_file = f"{base}_prompt.md"
        (run_dir / prompt_file).write_text(prompts.get(step.output_id, ""), encoding="utf-8")

        step_meta.append(
            {
                "step": step_num,
                "output_id": step.output_id,
                "display_name": step.display_name,
                "output_file": output_file,
                "prompt_file": prompt_file,
            }
        )

    run_info = {
        "kit": kit.name,
        "timestamp": datetime.now().strftime("%Y-%m-%d_%H-%M-%S"),
        "steps": step_meta,
    }
    (run_dir / "run.json").write_text(json.dumps(run_info, indent=2), encoding="utf-8")


def _extract_text(file_path: Path) -> str | None:
    """Extract text from common file types."""
    suffix = file_path.suffix.lower()
    if suffix in {".txt", ".md", ".csv", ".json"}:
        return file_path.read_text(encoding="utf-8", errors="ignore")
    if suffix == ".pdf":
        try:
            from pypdf import PdfReader

            reader = PdfReader(str(file_path))
            return "\n\n".join(page.extract_text() or "" for page in reader.pages)
        except Exception:
            return None
    if suffix == ".docx":
        try:
            from docx import Document

            doc = Document(str(file_path))
            return "\n\n".join(p.text for p in doc.paragraphs)
        except Exception:
            return None
    if suffix in {".xlsx", ".xls"}:
        try:
            from openpyxl import load_workbook

            wb = load_workbook(str(file_path), read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    parts.append(" | ".join(str(c) if c is not None else "" for c in row))
            return "\n".join(parts)
        except Exception:
            return None
    return None


def _cmd_batch(args: argparse.Namespace) -> None:
    """Run a kit over each file in a directory."""
    import copy

    from .mcp_client import close_mcp_servers, init_mcp_servers
    from .runner import run_kit_async

    input_dir = Path(args.input_dir)
    if not input_dir.exists():
        print(f"Error: Input directory not found: {input_dir}")
        sys.exit(1)

    supported_exts = {".docx", ".pdf", ".txt", ".md", ".csv", ".json", ".xlsx", ".xls"}
    input_files = sorted(
        f for f in input_dir.iterdir() if f.is_file() and f.suffix.lower() in supported_exts
    )
    if not input_files:
        print(f"No supported input files found in {input_dir}")
        sys.exit(1)

    print(f"Found {len(input_files)} file(s) to process in {input_dir}")

    fixed_dynamic: dict[str, str] = {}
    for spec in args.dynamic:
        if "=" not in spec:
            print(f"Error: Invalid --dynamic format: {spec}")
            sys.exit(1)
        res_id, value = spec.split("=", 1)
        fixed_dynamic[res_id] = value

    kit_path = resolve_kit_path(args.kit, args.base_path)
    local_mcp = kit_path / "mcp_servers.json"
    kit_config = str(local_mcp) if local_mcp.exists() else None

    try:
        template_kit = load_reasoning_kit(kit_path)
    except FileNotFoundError as e:
        print(f"Error: {e}")
        sys.exit(1)

    errors: list[tuple[str, str]] = []

    async def _run_one(file_path: Path, index: int) -> None:
        from datetime import datetime

        print(f"\n[{index}/{len(input_files)}] Processing: {file_path.name}")

        text = _extract_text(file_path)
        if text is None:
            print(f"  Warning: Could not extract text from {file_path.name}, skipping.")
            errors.append((file_path.name, "text extraction failed"))
            return

        print(f"  Extracted {len(text)} characters")

        kit = copy.deepcopy(template_kit)
        target = next((r for r in kit.resources.values() if r.resource_id == args.resource), None)
        if target is None:
            available = [r.resource_id for r in kit.resources.values()]
            print(f"Error: Resource '{args.resource}' not found in kit. Available: {available}")
            sys.exit(1)
        target.content = text

        for res_id, value in fixed_dynamic.items():
            res_match = next((r for r in kit.resources.values() if r.resource_id == res_id), None)
            if res_match:
                res_match.content = value
            else:
                print(f"  Warning: Resource '{res_id}' not found in kit, skipping.")

        collected_prompts: dict[str, str] = {}
        run_kwargs: dict = dict(verbose=args.verbose, collected_prompts=collected_prompts)
        if args.model:
            run_kwargs["model"] = args.model

        await init_mcp_servers(config_path="mcp_servers.json", kit_config_path=kit_config)
        try:
            outputs = await run_kit_async(kit, **run_kwargs)
        except Exception as e:
            print(f"  Error running kit on {file_path.name}: {e}")
            errors.append((file_path.name, str(e)))
            return
        finally:
            await close_mcp_servers()

        if not args.no_save:
            stem = file_path.stem
            timestamp = datetime.now().strftime("%Y-%m-%d_%H-%M-%S")
            run_path = Path(args.output_dir) / kit.name / stem / timestamp
            run_path.mkdir(parents=True, exist_ok=True)
            _write_run_files(kit, outputs, collected_prompts, run_path)
            print(f"  Saved to: {run_path}")

    async def _run_all() -> None:
        for i, file_path in enumerate(input_files, 1):
            await _run_one(file_path, i)

    asyncio.run(_run_all())

    print(f"\nBatch complete. Processed {len(input_files) - len(errors)}/{len(input_files)} files.")
    if errors:
        print(f"Failed ({len(errors)}):")
        for name, reason in errors:
            print(f"  - {name}: {reason}")
        sys.exit(1)


if __name__ == "__main__":
    main()